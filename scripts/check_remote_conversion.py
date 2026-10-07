"""隔离容器冒烟：API真实远程适配器与统一质量门，不调用模型。"""

import json
import sys
import tempfile
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def main():
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from layers import converter
    from layers.file_processing.runtime import get_file_processor_registry
    registry = get_file_processor_registry()
    deadline = time.monotonic() + 40
    while registry.engine_state("libreoffice").status.value != "ready":
        if time.monotonic() >= deadline:
            raise AssertionError("remote conversion engine not ready")
        time.sleep(.2)
    with tempfile.TemporaryDirectory(prefix="remote-conversion-smoke-") as temporary:
        root = Path(temporary)
        docx = root / "ordinary.docx"
        document = Document()
        document.add_paragraph("Ordinary remote DOCX conversion")
        document.save(docx)
        xlsx = root / "ordinary.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "Ordinary remote XLSX conversion"
        workbook.save(xlsx)
        pptx = root / "ordinary.pptx"
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[5]).shapes.title.text = "Ordinary remote PPTX conversion"
        presentation.save(pptx)
        records = []
        for source in (docx, xlsx, pptx):
            started = time.monotonic()
            result = converter.convert_file(str(source), "pdf")
            assert result.success and result.output_path, result.error_type
            try:
                records.append(dict(source_format=source.suffix.lstrip("."), success=True,
                    size_bytes=Path(result.output_path).stat().st_size,
                    seconds=time.monotonic() - started))
            finally:
                converter.cleanup_conversion_output(result.output_path)
    print(json.dumps(records))


if __name__ == "__main__":
    main()
