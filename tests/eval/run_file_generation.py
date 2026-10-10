"""三次虚构文件聊天；最多十二次模型HTTP尝试，无重试、联网或判卷。"""
import argparse
import base64
from io import BytesIO
import json
import logging
import os
from pathlib import Path
import secrets
import subprocess
import sys
import tempfile
import time
import uuid

SAMPLES = [
    {"id": "description", "mode": "expert", "request": "生成一份md文件：包含准备、执行、检查三个步骤，每步一句话，内容自行编写。"},
    {"id": "continuation", "mode": "expert", "filename": "范例.md",
     "text": "# 练习清单\n1. 准备：检查材料是否齐全；记录准备结果。\n",
     "request": "照附件第1项的格式写出第2、3、4项，每项两句话，生成包含原第1项及新增部分的完整新文件，不指定新格式。"},
    {"id": "typed_edit", "mode": "fast", "filename": "说明.txt",
     "text": "操作说明\n请先检查所有材料是否齐全，再认真检查工具是否完好，然后开始执行，并在结束后记录结果。\n",
     "request": "缩写这个文件中第二行，标题不变。"},
]
KNOWLEDGE_SAMPLES = [
    {"id":"knowledge_generation","mode":"expert","request":"按知识库里的退货规则写一份给客户的说明文件"},
    {"id":"general_template","mode":"expert","request":"写一份员工请假申请模板"},
]


def describe_plan(knowledge_fixture=False):
    if knowledge_fixture:
        return {"chat_requests":2,"estimated_model_calls":6,"model_call_hard_limit":6,
                "search_calls":0,"judge_calls":0,"retry_failed_requests":False}
    return {"chat_requests": 3, "estimated_model_calls": 9, "model_call_hard_limit": 12,
            "search_calls": 0, "judge_calls": 0, "retry_failed_requests": False}


def artifact_text(content, filename):
    """只读提取实际产物，不能假设未指定格式的生成必定返回文本。"""
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        from docx import Document
        document = Document(BytesIO(content))
        return "\n".join([p.text for p in document.paragraphs]
                         + [" | ".join(c.text for c in row.cells)
                            for table in document.tables for row in table.rows])
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages)
    return content.decode("utf-8-sig")


def record_artifact(item, content):
    """按真实SSE交付字段只读核对产物，不假设存在filename字段。"""
    item["actual_size_bytes"] = len(content)
    item["file_content"] = artifact_text(content, item["files"][0]["download_filename"])


def model_selection(response):
    """保留模型实际选的工具及决策理由，不用业务解析器把非法工具改成合法工具。"""
    result = {"selected_tools": [], "tool_decisions": [], "reasoning": None}
    data = response if isinstance(response, dict) else (
        response.model_dump() if hasattr(response, "model_dump") else {})
    choices = data.get("choices") or []
    message = choices[0].get("message") or {} if choices else {}
    for call in message.get("tool_calls") or []:
        function = call.get("function") or {}
        arguments = function.get("arguments") or {}
        invalid = False
        if isinstance(arguments, str):
            try:
                arguments = json.loads(arguments)
            except (ValueError, TypeError):
                arguments, invalid = {}, True
        if not isinstance(arguments, dict):
            arguments, invalid = {}, True
        name = function.get("name")
        result["selected_tools"].append(name)
        result["tool_decisions"].append({"tool_name": name, "reasoning": arguments.get("reasoning"),
            "arguments": arguments, "arguments_parse_failed": invalid})
    if result["tool_decisions"]:
        result["reasoning"] = result["tool_decisions"][0]["reasoning"]
    else:
        try:
            content = json.loads(message.get("content") or "null")
            if isinstance(content, dict):
                result["reasoning"] = content.get("reasoning")
        except (ValueError, TypeError):
            pass
    return result


def observe_model_selection(recorder):
    """只读挂到实际模型返回处；异常调用也有明确的空选择字段。"""
    original = recorder.original
    def observed(messages, tier="fast", **kwargs):
        item = recorder.current_call.get()
        if item is not None:
            item.update({"selected_tools": [], "tool_decisions": [], "reasoning": None})
        response = original(messages, tier=tier, **kwargs)
        if item is not None:
            item.update(model_selection(response))
        return response
    recorder.original = observed


def generation_route_verified(item):
    return (item.get("selected_tools") == ["generate_file"]
            and not item["tool_decisions"][0]["arguments_parse_failed"])


def run(args):
    repo = Path(__file__).resolve().parents[2]
    sys.path.insert(0, str(repo))
    from tests.eval.run_eval import CallRecorder, EvalStopped, is_project_venv, snapshot_data, source_revision, write_json
    knowledge_fixture = getattr(args, "knowledge_fixture", False)
    plan = describe_plan(knowledge_fixture)
    print("PREPARE " + json.dumps(plan, ensure_ascii=False), flush=True)
    if args.prepare_only:
        return
    if not is_project_venv(repo) or not 1 <= args.max_calls <= plan["model_call_hard_limit"]:
        raise EvalStopped("Use project .venv and stay within scenario hard limit")
    sys.dont_write_bytecode = True
    output = (repo / "backups/eval" / args.output).resolve()
    if output.parent != (repo / "backups/eval").resolve():
        raise EvalStopped("Output must be directly under ignored backups/eval")
    output.mkdir(exist_ok=False, parents=True)
    before = snapshot_data(repo / "data")
    write_json(output / "default_data_before.json", before)
    work = Path(tempfile.mkdtemp(prefix="zhitian-eval-runtime-")).resolve()
    private_temp = work / "tmp"
    private_temp.mkdir()
    os.environ["TEMP"] = os.environ["TMP"] = str(private_temp)
    tempfile.tempdir = str(private_temp)
    os.environ["JWT_SECRET_KEY"] = secrets.token_urlsafe(48)
    os.environ["ENTERPRISE_PASSWORD_SEED"] = secrets.token_urlsafe(48)
    os.environ["PERSONAL_DEEPSEEK_KEY_ENCRYPTION_KEY"] = base64.urlsafe_b64encode(secrets.token_bytes(32)).decode()
    os.environ["SCHEDULED_BACKUP_ENABLED"] = "false"
    import config
    config.BASE_DIR = str(work)
    config.HISTORY_DB_PATH = str(work / "data/history.db")
    config.VECTORDB_PATH = str(work / "data/vectordb")
    config.SCHEDULED_BACKUP_PATH = str(work / "backups")

    def guard(event, values):
        if event != "open" or not isinstance(values[0], (str, bytes, os.PathLike)):
            return
        filename, mode, flags = values
        if os.fsdecode(filename).lower() == os.devnull.lower():
            return
        writing = (isinstance(mode, str) and any(x in mode for x in "wax+")) or (isinstance(flags, int) and flags & (os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_TRUNC))
        if writing:
            path = Path(os.fsdecode(filename)).resolve()
            if not any(path == root or root in path.parents for root in (work, output)):
                raise EvalStopped("Write outside isolated runtime/results blocked")
    sys.addaudithook(guard)
    from fastapi.testclient import TestClient
    import main
    from layers import auth, api_quota, enterprise_password, execution, llm_provider, memory
    from layers.file_processing.runtime import get_file_processor_registry
    recorder = CallRecorder(llm_provider, output, args.max_calls, no_judge=True)
    observe_model_selection(recorder)

    def prevent_retry(request):
        if request.method == "POST" and "/chat/completions" in request.url.path:
            item = recorder.current_call.get()
            if item and item["attempts"]:
                recorder.stop_reason = "retry_forbidden"
                raise EvalStopped("Failed/timeout model calls must not be retried")
    llm_provider._get_shared_http_client().event_hooks["request"].append(prevent_retry)
    recorder.install_http_hooks()

    def recorded_call(messages, tier="fast", **kwargs):
        if recorder.stop_reason:
            raise EvalStopped("Smoke already stopped")
        stage = str(getattr(kwargs.get("stage"), "value", kwargs.get("stage")))
        if knowledge_fixture and stage not in {
                "intent_classification", "direct_chat_reasoning", "memory_importance"}:
            recorder.stop_reason = "unexpected_model_stage:" + stage
            raise EvalStopped("Generation smoke must not run rerank, reflection, planning or search models")
        # 只限制本次测量的重试，不改变应用默认配置、预算或模型参数。
        kwargs["retry_timeouts"] = False
        response = recorder.call(messages, tier=tier, **kwargs)
        if knowledge_fixture and stage == "intent_classification":
            item = next(r for r in reversed(recorder.records)
                        if r["round"] == recorder.current and r["stage"] == stage)
            write_json(output / "model_calls.json", recorder.records)
            if not generation_route_verified(item):
                recorder.stop_reason = "classification_not_generate_file"
                print("ROUTE_STOP " + json.dumps({"selected_tools":item["selected_tools"],
                    "tool_decisions":item["tool_decisions"]},ensure_ascii=False),flush=True)
                raise EvalStopped("Classifier did not select generate_file; stop without retry")
        return response
    llm_provider.chat_completion = recorded_call

    def forbidden(*_a, **_k):
        raise EvalStopped("File task must not search or write long-term memory")
    local_searches = []
    original_search = execution._search_documents
    def local_only_search(*a, **k):
        if recorder.stop_reason:
            raise EvalStopped("Smoke already stopped before knowledge search")
        if k.get("rerank_enabled") is not False or k.get("generate_answer") is not False:
            raise EvalStopped("Generation knowledge search must be local only")
        local_searches.append({"round":recorder.current,"rerank_enabled":False,"generate_answer":False})
        return original_search(*a, **k)
    execution._search_documents = local_only_search
    execution._search_web = forbidden
    original_memory_save = memory.maybe_save_to_vector
    def user_memory_only(session, role, content, *a, **k):
        if role != "user":
            raise EvalStopped("Generated content must not enter long-term memory")
        return original_memory_save(session, role, content, *a, **k)
    memory.maybe_save_to_vector = user_memory_only
    completed = []
    try:
        if not config.DEEPSEEK_API_KEY:
            raise EvalStopped("Missing model credential")
        get_file_processor_registry().probe_sync("native_text")
        get_file_processor_registry().probe_sync("document_text")
        password = secrets.token_urlsafe(20)
        user = auth.register_user("isolated-file-fixture", password, "customer")
        api_quota.authorize_enterprise_source(user["user_id"], enterprise_password.get_current_enterprise_password())
        headers = {"Authorization": "Bearer " + auth.login_user("isolated-file-fixture", password, "customer")}
        client = TestClient(main.app)
        if knowledge_fixture:
            prepare_knowledge(client, auth, password, output)
        for sample in KNOWLEDGE_SAMPLES if knowledge_fixture else SAMPLES:
            if recorder.stop_reason:
                break
            recorder.current = sample["id"]
            count_before = recorder.model_count
            session = uuid.uuid4().hex
            payload = {"session_id": session, "message": sample["request"], "mode": sample["mode"], "attachment_page_id": "fixture-page"}
            raw = sample.get("text", "").encode("utf-8")
            if sample.get("filename"):
                upload = client.post("/chat/attachments", headers=headers, data={"session_id":session, "page_id":"fixture-page"},
                    files={"file":(sample["filename"], raw, "text/plain")})
                upload.raise_for_status()
                payload["attachment_ids"] = [upload.json()["attachment_id"]]
            start = time.perf_counter()
            if sample["id"] == "typed_edit":
                response = client.post("/chat/stream/originals", headers=headers,
                    data={"payload":json.dumps(payload), "original_ids":json.dumps(payload["attachment_ids"])},
                    files=[("files",(sample["filename"],raw,"text/plain"))])
            else:
                response = client.post("/chat/stream", headers=headers, json=payload)
            response.raise_for_status()
            elapsed = time.perf_counter()-start
            events = [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith("data: ")]
            files = [event for event in events if event.get("type") == "file"]
            item = {"id":sample["id"], "mode":sample["mode"], "elapsed_seconds":elapsed,
                "model_calls":recorder.model_count-count_before, "files":files,
                "status":next((e for e in events if e.get("type") == "request_status"),None),
                "answer":"".join(e.get("chunk","") for e in events if e.get("chunk") != "[DONE]")}
            item["knowledge_searches"] = sum(s["round"] == sample["id"] for s in local_searches)
            item["classification"] = next(({"selected_tools":r["selected_tools"],
                "tool_decisions":r["tool_decisions"],"reasoning":r["reasoning"]}
                for r in recorder.records if r["round"] == sample["id"]
                and r["stage"] == "intent_classification"),None)
            item["source_policy"] = next((e for e in events if e.get("type") == "source_policy"),None)
            item["citations"] = next((e["citations"] for e in events if e.get("type") == "citations"),[])
            completed.append(item)
            # 下载或正文提取失败时，仍保留已经收到的状态、引用与文件卡。
            write_json(output / "results.json",completed)
            if files:
                artifact = client.get("/files/"+files[0]["file_id"],headers=headers)
                artifact.raise_for_status()
                record_artifact(item, artifact.content)
                client.post("/files/"+files[0]["file_id"]+"/receipt",headers=headers).raise_for_status()
            write_json(output / "results.json",completed)
            print("RESULT " + json.dumps(item,ensure_ascii=False),flush=True)
            if (not files or recorder.stop_reason
                or any(r["error_type"] for r in recorder.records if r["round"]==recorder.current)):
                break
    finally:
        llm_provider.close_resources()
        logging.shutdown()
        after = snapshot_data(repo / "data")
        write_json(output / "default_data_after.json",after)
        write_json(output / "model_calls.json",recorder.records)
        write_json(output / "local_searches.json",local_searches)
        write_json(output / "metadata.json",{"source_revision":source_revision(repo),
            "model_calls":recorder.model_count,"search_calls":recorder.web_count,
            "default_data_unchanged":before==after,"runtime_directory":str(work),"stop_reason":recorder.stop_reason})
        print("CALLS " + str(recorder.model_count) + " DATA_UNCHANGED " + str(before==after),flush=True)
        if before != after:
            raise EvalStopped("Default data changed")


def prepare_knowledge(client, auth, password, output):
    """现场虚构MD，真实上传与审核，零模型调用。"""
    from tests.eval.run_eval import EvalStopped, write_json
    def request(method, route, **kwargs):
        response = client.request(method, route, **kwargs)
        response.raise_for_status()
        return response.json()
    roles = {}
    for role in ("developer","employee","reviewer"):
        name = "knowledge_fixture_" + role
        auth.register_user(name,password,role)
        roles[role] = {"Authorization":"Bearer " + auth.login_user(name,password,role)}
    organization = request("POST","/developer/organizations",headers=roles["developer"],
        json={"name":"隔离评估资料","content":"本地虚构规则，仅用于隔离验证"})
    for role in ("employee","reviewer"):
        application = request("POST",f"/organizations/{organization['id']}/join-request",headers=roles[role])
        request("POST",f"/developer/org-membership-requests/{application['id']}/approve",headers=roles["developer"])
    raw = "# 退货规则\n\n退货期限为7天。申请退货时必须保留发票。\n".encode("utf-8")
    accepted = request("POST","/documents/upload",headers=roles["employee"],
        data={"organization_id":str(organization["id"])},files={"file":("退货规则.md",raw,"text/markdown")})
    deadline = time.monotonic() + 180
    while True:
        task = request("GET","/tasks/"+accepted["task_id"],headers=roles["employee"])
        if task.get("status") in {"done","failed"} or time.monotonic() >= deadline:
            break
        time.sleep(.1)
    if task.get("status") != "done":
        raise EvalStopped("Isolated knowledge upload failed")
    verified = request("POST","/approve/"+accepted["doc_id"],headers=roles["reviewer"])
    if verified.get("status") != "verified":
        raise EvalStopped("Isolated knowledge review failed")
    write_json(output/"knowledge_fixture.json",{"doc_id":accepted["doc_id"],"status":"verified","chunks":accepted["chunks"]})


if __name__ == "__main__":
    parser=argparse.ArgumentParser()
    parser.add_argument("--prepare-only",action="store_true")
    parser.add_argument("--max-calls",type=int,default=12)
    parser.add_argument("--output",default="file-generation-20261010")
    parser.add_argument("--knowledge-fixture",action="store_true",help="Two expert requests, real isolated knowledge upload, max six model calls")
    parser.add_argument("--worker",action="store_true",help=argparse.SUPPRESS)
    args=parser.parse_args()
    if args.prepare_only or args.worker:
        run(args)
    else:
        process=subprocess.run([sys.executable,"-B",str(Path(__file__).resolve()),"--worker",
            "--max-calls",str(args.max_calls),"--output",args.output]
            + (["--knowledge-fixture"] if args.knowledge_fixture else []))
        repo=Path(__file__).resolve().parents[2]
        sys.path.insert(0,str(repo))
        from tests.eval.run_eval import cleanup_runtime,write_json
        metadata_path=repo/"backups/eval"/args.output/"metadata.json"
        if metadata_path.is_file():
            metadata=json.loads(metadata_path.read_text(encoding="utf-8"))
            cleanup_runtime(metadata["runtime_directory"])
            metadata["runtime_removed"]=not Path(metadata["runtime_directory"]).exists()
            write_json(metadata_path,metadata)
        raise SystemExit(process.returncode)
