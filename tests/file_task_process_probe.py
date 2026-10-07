"""仅隔离Linux容器使用：暂停真实soffice模拟卡住，验证超时终止与资源释放。"""

import json
import os
from pathlib import Path
import signal
import sys
import tempfile
import time


def processes():
    result = []
    for path in Path("/proc").iterdir():
        if not path.name.isdigit():
            continue
        try:
            fields = (path / "stat").read_text().rsplit(")", 1)[1].split()
            command = (path / "cmdline").read_bytes().replace(b"\0", b" ").decode(errors="replace")
            result.append(dict(pid=int(path.name), parent=int(fields[1]), group=int(fields[2]),
                               state=fields[0], executable=command.split(" ")[0]))
        except (OSError, IndexError):
            pass
    return result


def main():
    if sys.platform != "linux":
        raise SystemExit("isolated_linux_container_required")
    from utils import logger
    logger._configured = True
    import config
    with tempfile.TemporaryDirectory(prefix="zhitian-process-probe-") as directory:
        config.BASE_DIR = directory
        config.HISTORY_DB_PATH = str(Path(directory) / "data/history.db")
        config.VECTORDB_PATH = str(Path(directory) / "data/vectordb")
        from docx import Document
        from layers import converter, heavy_task_limits, resource_admission
        from layers.file_processing import runner
        from layers.file_processing.runtime import get_file_processor_registry
        from layers.file_processing.models import EngineStatus
        registry = get_file_processor_registry()
        registry._states["libreoffice"].status = EngineStatus.READY
        source = Path(directory) / "probe.docx"
        document = Document()
        for index in range(2000):
            document.add_paragraph("Controlled local conversion probe %d" % index)
        document.save(source)
        baseline = resource_admission.read_cgroup_memory()[0]
        original = converter.run_process
        during = []
        snapshots = []
        stopped = False
        def run(command, workspace, scope):
            def poll():
                nonlocal stopped, during
                snapshot = resource_admission.read_cgroup_memory()[0]
                snapshots.append(snapshot)
                active = processes()
                owned = [item for item in active if any(item["group"] == p.pid for p in workspace.processes)]
                if not stopped and any("soffice.bin" in item["executable"] for item in owned):
                    during = owned
                    os.killpg(next(iter(workspace.processes)).pid, signal.SIGSTOP)
                    stopped = True
            return runner.run_process(command, workspace, scope, on_poll=poll)
        converter.run_process = run
        slots_before = heavy_task_limits.slots_in_use()
        reserved_before = resource_admission.reserved_bytes()
        before_processes = [item for item in processes()
                            if any(name in item["executable"] for name in ("oosplash", "soffice.bin"))]
        assert not before_processes
        started = time.monotonic()
        with runner.task_scope(3), heavy_task_limits.occupy_slot():
            result = converter.convert_file(str(source), "pdf")
        after = resource_admission.read_cgroup_memory()[0]
        leftovers = [item for item in processes() if item["pid"] in {item["pid"] for item in during}]
        assert stopped, "real_soffice_not_observed"
        assert result.status == converter.ConversionStatus.TIMEOUT
        assert not leftovers, leftovers
        assert heavy_task_limits.slots_in_use() == slots_before
        assert resource_admission.reserved_bytes() == reserved_before
        assert converter._conversion_lock.acquire(blocking=False)
        converter._conversion_lock.release()
        assert not list(runner.task_root().glob("task_*"))
        converter.run_process = original
        print(json.dumps(dict(result=result.status.value, elapsed_seconds=time.monotonic() - started,
            processes_before=before_processes, processes_during=during, processes_after=leftovers,
            slots_before=slots_before, slots_after=heavy_task_limits.slots_in_use(),
            reserved_before=reserved_before, reserved_after=resource_admission.reserved_bytes(),
            memory_before=baseline.current_bytes, memory_peak=max(item.current_bytes for item in snapshots),
            memory_after=after.current_bytes,
            adjusted_before=baseline.limit_bytes - baseline.available_bytes,
            adjusted_after=after.limit_bytes - after.available_bytes), ensure_ascii=False))

        # 同一隔离容器确认正常转换，及独立PDF解析进程的额外内存。
        document = Document()
        document.add_paragraph("Normal local conversion after timeout")
        document.save(source)
        with heavy_task_limits.occupy_slot():
            normal = converter.convert_file(str(source), "pdf")
        assert normal.success, normal.error_type
        print(json.dumps(dict(normal_conversion=normal.status.value,
            stages=[event.stage for event in normal.progress_events])))
        converter.cleanup_conversion_output(normal.output_path)
        import fitz
        from layers.file_processing.pdf import pdf_processor
        from layers.file_processing.models import FileProcessingRequest
        pdf_path = Path(directory) / "pages.pdf"
        with fitz.open() as pdf:
            for index in range(150):
                pdf.new_page().insert_text((72, 72), "Local page %d" % index)
            pdf.save(pdf_path)
        worker_before = resource_admission.read_cgroup_memory()[0]
        worker_memory, worker_rss = [], []
        def measured_run(command, workspace, scope, on_poll=None):
            def poll():
                if on_poll:
                    on_poll()
                worker_memory.append(resource_admission.read_cgroup_memory()[0])
                for process in workspace.processes:
                    try:
                        status = Path("/proc/%s/status" % process.pid).read_text()
                        line = next(line for line in status.splitlines() if line.startswith("VmRSS:"))
                        worker_rss.append(int(line.split()[1]) * 1024)
                    except (OSError, StopIteration):
                        pass
            return original_run(command, workspace, scope, on_poll=poll)
        original_run = runner.run_process
        runner.run_process = measured_run
        try:
            parsed = pdf_processor.execute_task(FileProcessingRequest(task_type="extract", source_format="pdf",
                source_paths=[str(pdf_path)]))
        finally:
            runner.run_process = original_run
        assert parsed.success and parsed.page_count == 150
        assert not list(runner.task_root().glob("task_*"))
        worker_after = resource_admission.read_cgroup_memory()[0]
        print(json.dumps(dict(pdf_pages=parsed.page_count, worker_peak_rss=max(worker_rss),
            memory_before=worker_before.current_bytes,
            memory_peak=max(item.current_bytes for item in worker_memory),
            memory_after=worker_after.current_bytes,
            adjusted_peak_increment=max(item.limit_bytes - item.available_bytes for item in worker_memory)
                - (worker_before.limit_bytes - worker_before.available_bytes))))


if __name__ == "__main__":
    main()
