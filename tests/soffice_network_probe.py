# -*- coding: utf-8 -*-
"""镜像构建期真实回环探针：验证 soffice 子进程不外连且正常转换不变。"""

import hashlib
import json
import os
import re
import subprocess
import sys
import tempfile
import threading
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

# 此脚本由 Dockerfile 在镜像内直接执行，不加载 tests/conftest.py；测试凭据
# 仅满足 config 的导入校验，不连接模型服务，也不触碰运行时数据卷。
os.environ["JWT_SECRET_KEY"] = "test-only-jwt-secret-at-least-32-bytes-2026"
os.environ["ENTERPRISE_PASSWORD_SEED"] = "test-only-enterprise-password-seed-not-for-production"
os.environ["DEEPSEEK_API_KEY"] = "test-only-deepseek-key-not-for-production"
os.environ["PERSONAL_DEEPSEEK_KEY_ENCRYPTION_KEY"] = (
    "dGVzdC1vbmx5LXBlcnNvbmFsLWtleS1zZWNyZXQhISE="
)
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402

# 构建期没有业务数据卷，不能让探针初始化项目文件日志并写入 /app/data。
# 仅关闭此探针进程的日志初始化；转换链路、网络隔离和探针断言不变。
from utils import logger as project_logger  # noqa: E402

project_logger._configured = True

from layers import converter, document_loader  # noqa: E402


def _measure(path: Path) -> dict:
    assert path.is_file(), "LibreOffice 未产出 PDF"
    text = document_loader.load_document(str(path))
    assert text and not text.startswith("错误："), "PDF 不能重新解析"
    return {
        "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
        "chunk_count": len(document_loader.chunk_text(text)),
    }


def _plain_conversion(source: Path, root: Path, label: str) -> Path:
    output_dir = root / label
    output_dir.mkdir()
    profile_url = (root / (label + "-profile")).as_uri()
    completed = subprocess.run(
        [
            "/usr/bin/soffice",
            "-env:UserInstallation=" + profile_url,
            "--headless",
            "--convert-to", "pdf",
            "--outdir", str(output_dir), str(source),
        ],
        capture_output=True,
        check=False,
        timeout=30,
    )
    assert completed.returncode == 0, "未隔离基线转换失败"
    return output_dir / (source.stem + ".pdf")


def _linked_docx(source: Path, destination: Path, target: str) -> None:
    """仅把有效 DOCX 内的图片关系改成容器回环外链。"""
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(destination, "w") as linked:
        for member in original.infolist():
            payload = original.read(member.filename)
            if member.filename == "word/_rels/document.xml.rels":
                payload, count = re.subn(
                    rb'Target="media/[^\"]+"',
                    b'Target="' + target.encode("ascii") + b'" TargetMode="External"',
                    payload,
                    count=1,
                )
                assert count == 1, "图片关系未被替换"
            if member.filename == "word/document.xml":
                payload, count = re.subn(
                    rb'r:embed="([^"]+)"', rb'r:link="\1"', payload, count=1
                )
                assert count == 1, "正文图片引用未被替换"
            linked.writestr(member, payload)


def run_probe() -> dict:
    # 反证隔离确已安装，而非恰好这份文档没有触发远程加载。
    sandbox = str(Path(converter.__file__).with_name("soffice_sandbox.py"))
    network_socket = subprocess.run(
        [
            sys.executable, sandbox, sys.executable, "-c",
            "import socket; socket.socket(socket.AF_INET, socket.SOCK_STREAM)",
        ],
        capture_output=True, check=False, timeout=10,
    )
    assert network_socket.returncode != 0, "隔离进程仍可创建网络 socket"
    unix_socket = subprocess.run(
        [
            sys.executable, sandbox, sys.executable, "-c",
            "import socket; socket.socket(socket.AF_UNIX, socket.SOCK_STREAM)",
        ],
        capture_output=True, check=False, timeout=10,
    )
    assert unix_socket.returncode == 0, "隔离错误地阻断了 LibreOffice 所需的 Unix socket"
    hits = []

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            hits.append(self.path)
            self.send_response(200)
            self.send_header("Content-Type", "image/png")
            self.send_header("Content-Length", str(len(image_bytes)))
            self.end_headers()
            self.wfile.write(image_bytes)

        def log_message(self, *_args):
            pass

    with tempfile.TemporaryDirectory(prefix="soffice-network-test-") as temporary:
        root = Path(temporary)
        image_path = root / "image.png"
        Image.new("RGB", (1, 1), (255, 0, 0)).save(image_path)
        image_bytes = image_path.read_bytes()

        docx = root / "normal.docx"
        document = Document()
        document.add_paragraph("普通文档：中文正文测试。")
        document.add_picture(str(image_path))
        document.save(docx)
        xlsx = root / "normal.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "普通表格：中文内容"
        workbook.save(xlsx)
        pptx = root / "normal.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "普通幻灯片：中文标题"
        presentation.save(pptx)

        normal = {}
        for source in (docx, xlsx, pptx):
            baseline = _measure(_plain_conversion(source, root, source.suffix[1:] + "-old"))
            secured = converter.convert_file(str(source), "pdf")
            assert secured.success and secured.output_path, "隔离后普通文档转换失败"
            after = _measure(Path(secured.output_path))
            assert after == baseline, "隔离改变了普通文档的正文或切片数"
            normal[source.suffix] = after
            converter.cleanup_conversion_output(secured.output_path)

        with ThreadingHTTPServer(("127.0.0.1", 0), Handler) as server:
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            target = "http://127.0.0.1:%s/probe.png" % server.server_port
            linked = root / "linked.docx"
            _linked_docx(docx, linked, target)
            _measure(_plain_conversion(linked, root, "linked-old"))
            before = len(hits)
            assert before > 0, "对照转换没有触发回环请求，探针失效"
            secured = converter.convert_file(str(linked), "pdf")
            assert secured.success and secured.output_path, "隔离后外链文档转换失败"
            _measure(Path(secured.output_path))
            after = len(hits) - before
            converter.cleanup_conversion_output(secured.output_path)
            server.shutdown()
            assert after == 0, "隔离后的转换仍发出了回环网络请求"
        return {
            "linked_before_requests": before,
            "linked_after_requests": after,
            "normal": normal,
        }


if __name__ == "__main__":
    print(json.dumps(run_probe(), ensure_ascii=True))
