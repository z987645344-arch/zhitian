"""DOCX tables use real OOXML fixtures; no model, network or default data."""
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement

from layers import document_loader as loader, document_sections as sections


def read(doc, tmp_path):
    path = tmp_path / "table.docx"
    doc.save(path)
    return path, loader.load_document(str(path))


def table(doc, rows):
    result = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for row, values in zip(result.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
    return result


def test_body_order_headers_blank_rows_and_cells(tmp_path):
    doc = Document()
    doc.add_paragraph("前文")
    table(doc, [["名称", "价格", "备注"], ["甲", "17元", ""],
                ["", "", ""], ["", "9元", "后项"]])
    doc.add_paragraph("中间正文")
    table(doc, [["编号", "值"], ["B", "二"]])
    doc.add_paragraph("尾文")
    _, text = read(doc, tmp_path)
    assert text.splitlines() == ["前文", "名称 | 价格 | 备注", "甲 | 17元 |",
                                "| 9元 | 后项", "中间正文", "编号 | 值", "B | 二", "尾文"]


def test_horizontal_and_vertical_merge_do_not_repeat_origin_text(tmp_path):
    doc = Document()
    grid = table(doc, [["列一", "列二", "列三"], ["横合并", "", "其他"],
                       ["纵合并", "相同", "相同"], ["", "独立", "后续"]])
    grid.cell(1, 0).merge(grid.cell(1, 1))
    grid.cell(2, 0).merge(grid.cell(3, 0))
    _, text = read(doc, tmp_path)
    assert text.count("横合并") == text.count("纵合并") == 1
    assert text.count("相同") == 2  # Different cells are not deduplicated by value.
    assert "横合并 | 其他" in text
    assert "| 独立 | 后续" in text


def test_nested_table_and_cell_paragraphs_stay_in_place(tmp_path):
    doc = Document()
    outer = table(doc, [["外列", "说明"], ["", "右格"]])
    cell = outer.cell(1, 0)
    cell.paragraphs[0].text = "嵌套之前"
    nested = cell.add_table(rows=2, cols=2)
    nested.cell(0, 0).text, nested.cell(0, 1).text = "内列", "内值"
    nested.cell(1, 0).text, nested.cell(1, 1).text = "唯一嵌套事实", "23元"
    cell.add_paragraph("嵌套之后")
    _, text = read(doc, tmp_path)
    assert "嵌套之前 ; 内列 | 内值 ; 唯一嵌套事实 | 23元 ; 嵌套之后 | 右格" in text
    assert text.count("唯一嵌套事实") == 1


def test_long_table_repeats_header_and_preserves_all_rows_and_sections(tmp_path):
    doc = Document()
    doc.add_heading("文档标题", level=0)
    doc.add_heading("大节", level=1)
    doc.add_heading("费用表", level=2)
    header = "服务编号 | 费用"
    grid = table(doc, [["服务编号", "费用"]] + [[f"条目{i:03}", f"{i}元"] for i in range(40)])
    grid.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_heading("下一节", level=1)
    doc.add_paragraph("下一节的独立正文")
    path, text = read(doc, tmp_path)
    chunks = loader.chunk_text(text, chunk_size=60)
    fields = sections.chunk_section_paths(text, chunks, source_path=str(path))
    assert len(chunks) > 8
    assert text.count(header) == 1
    for index in range(40):
        containing = [(chunk, field) for chunk, field in zip(chunks, fields) if f"条目{index:03}" in chunk]
        assert len(containing) == 1
        chunk, field = containing[0]
        assert header in chunk
        assert ("文档标题", "大节", "费用表") in sections.read_section_paths(field)
    assert chunks[-1].endswith("下一节的独立正文")
    assert ("文档标题", "下一节") in sections.read_section_paths(fields[-1])
    assert all(fields)


def test_heading_offsets_include_tables_and_cross_section_chunks(tmp_path):
    doc = Document()
    doc.add_heading("总标题", 0)
    doc.add_heading("甲节", 1)
    table(doc, [["类型", "价格"], ["甲型", "17元"]])
    doc.add_heading("乙节", 1)
    doc.add_heading("乙小节", 2)
    table(doc, [["另一列", "值"], ["乙型", "9元"]])
    path, text = read(doc, tmp_path)
    chunks = [text]
    fields = sections.chunk_section_paths(text, chunks, source_path=str(path))
    assert sections.read_section_paths(fields[0]) == {("总标题", "甲节"), ("总标题", "乙节"),
                                                    ("总标题", "乙节", "乙小节")}
    assert list(loader.iter_docx_text_blocks(Document(path))) == [
        ("总标题", "Title"), ("甲节", "Heading1"), ("类型 | 价格", ""), ("甲型 | 17元", ""),
        ("乙节", "Heading1"), ("乙小节", "Heading2"), ("另一列 | 值", ""), ("乙型 | 9元", ""),
    ]


def test_cell_heading_does_not_create_a_body_section(tmp_path):
    doc = Document()
    doc.add_heading("标题", 0)
    doc.add_heading("实际小节", 1)
    grid = table(doc, [["列", "值"], ["单元格标题", "正文"]])
    grid.cell(1, 0).paragraphs[0].style = "Heading 1"
    path, text = read(doc, tmp_path)
    fields = sections.chunk_section_paths(text, loader.chunk_text(text), source_path=str(path))
    assert sections.read_section_paths(fields[0]) == {("标题", "实际小节")}


def test_long_cell_splitting_keeps_text_and_repeated_headers(tmp_path):
    doc = Document()
    table(doc, [["编号", "说明"], ["A", "长" * 240]])
    _, text = read(doc, tmp_path)
    chunks = loader.chunk_text(text, chunk_size=40)
    assert len(chunks) > 4
    assert all(chunk.startswith("编号 | 说明") for chunk in chunks)
    assert sum(chunk.count("长") for chunk in chunks) == 240


def test_oversized_header_is_retained_without_unbounded_duplication(tmp_path):
    doc = Document()
    table(doc, [["列" * 100, "其他列"], ["值" * 100, "其他值"]])
    _, text = read(doc, tmp_path)
    chunks = loader.chunk_text(text, chunk_size=40)
    assert sum(chunk.count("列") for chunk in chunks) == 101
    assert sum(chunk.count("值") for chunk in chunks) == 101


def test_multiple_explicit_header_rows_repeat_together(tmp_path):
    doc = Document()
    grid = table(doc, [["组名", "范围"], ["编号", "价格"]] + [[str(i), "金额"] for i in range(20)])
    for row in grid.rows[:2]:
        row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    _, text = read(doc, tmp_path)
    chunks = loader.chunk_text(text, chunk_size=50)
    assert all("组名 | 范围\n编号 | 价格" in chunk for chunk in chunks)


def test_paragraph_only_docx_and_plain_text_chunking_unchanged(tmp_path):
    doc = Document()
    doc.add_paragraph("第一段")
    doc.add_paragraph("")
    doc.add_paragraph("第二段。" * 80)
    _, text = read(doc, tmp_path)
    assert text == "第一段\n" + "第二段。" * 80
    assert loader.chunk_text(text) == loader.chunk_text(str(text))
    for plain in ["# 标题\n甲 | 乙\n丙" * 30, "一般纯文本。" * 100]:
        assert all(type(chunk) is str for chunk in loader.chunk_text(plain))


def test_actual_chroma_store_keeps_repeated_headers_and_optional_sections(tmp_path):
    from layers import memory
    doc = Document()
    doc.add_heading("数据表", 0)
    doc.add_heading("定价", 1)
    table(doc, [["项目", "费用"]] + [[f"项目{i}", f"{i}元"] for i in range(30)])
    path, text = read(doc, tmp_path)
    chunks = loader.chunk_text(text, chunk_size=50)
    fields = sections.chunk_section_paths(text, chunks, source_path=str(path))
    assert memory.save_document("table.docx", chunks, "docx-table-test", chunk_section_paths=fields) == len(chunks)
    rows = memory._get_document_collection().get(where={"doc_id": "docx-table-test"},
                                                include=["documents", "metadatas"])
    records = sorted(zip(rows["documents"], rows["metadatas"]), key=lambda row: row[1]["chunk_index"])
    assert [body for body, _ in records] == chunks
    assert all(sections.read_section_paths(metadata.get("section_paths")) == {("数据表", "定价")}
               for _, metadata in records)


@pytest.mark.parametrize("question_id", ["S21", "S22"])
def test_eval_table_only_quote_is_actually_extracted_and_chunked(question_id):
    import json
    root = Path(__file__).parent / "eval"
    dataset = json.loads((root / "questions.json").read_text(encoding="utf-8"))
    question = next(item for item in dataset["questions"] if item["id"] == question_id)
    evidence = question["supporting_evidence"][0]
    text = loader.load_document(str(root / "corpus" / evidence["source"]))
    assert evidence["quote"] in text
    assert any(evidence["quote"] in chunk for chunk in loader.chunk_text(text))
