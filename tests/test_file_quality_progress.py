"""统一结果、质量门、真实阶段/计数；零付费。"""

import asyncio
import json
from pathlib import Path

import fitz
import pytest
from docx import Document
from openpyxl import load_workbook
from pydantic import ValidationError

import config
import main
from layers import converter, execution, task_store
from layers.file_processing.base import FileProcessor
from layers.file_processing.degradation import DEGRADATIONS, OFFICE_TO_MARKDOWN
from layers.file_processing.models import (
    FileArtifact, FileProcessingRequest, FileProcessingResult, FileProcessingStatus,
    FileTaskProgress, QualityProfile, QualityCheckResult,
)
from layers.file_processing.pdf import pdf_processor
from layers.file_processing.quality import FileQualityChecker
from layers.file_processing.runtime import get_file_processor_registry


def _pdf(path, pages=2, table=False):
    with fitz.open() as document:
        for _ in range(pages):
            page = document.new_page()
            if table:
                for x in (50, 200, 350):
                    page.draw_line((x, 50), (x, 150))
                for y in (50, 100, 150):
                    page.draw_line((50, y), (350, y))
                for x, y, text in ((70, 80, "Item"), (220, 80, "Value"),
                                   (70, 130, "Model"), (220, 130, "42")):
                    page.insert_text((x, y), text)
            else:
                page.insert_text((72, 72), "Normal document text")
        document.save(path)


@pytest.mark.parametrize("status,success", [("SUCCESS", True), ("FAILED", False),
                                          ("TIMEOUT", False), ("CANCELLED", False)])
def test_explicit_outcomes(status, success):
    result = FileProcessingResult(success=success, status=status)
    assert result.success == success
    assert result.degradation_code == ""


def test_degradation_requires_prior_registration_and_exact_notice():
    for code, notice in (("unknown", "anything"), (OFFICE_TO_MARKDOWN, "")):
        with pytest.raises(ValidationError):
            FileProcessingResult(success=True, status="DEGRADED", degradation_code=code, user_notice=notice)
    result = FileProcessingResult(success=True, status="DEGRADED", degradation_code=OFFICE_TO_MARKDOWN,
                                  user_notice=DEGRADATIONS[OFFICE_TO_MARKDOWN])
    assert result.status == FileProcessingStatus.DEGRADED
    with pytest.raises(ValidationError):
        FileProcessingResult(success=True, status="FAILED")


def test_docx_quality_counts_table_text_and_rejects_corrupted_table(tmp_path):
    source = tmp_path / "table.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Only in table"
    table.cell(1, 1).text = "42"
    document.save(source)
    artifact = FileArtifact(output_path=str(source), file_format="docx",
                            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    checked = FileQualityChecker().validate(artifact, QualityProfile.DOCX, minimum_paragraphs=1)
    assert checked.passed and artifact.paragraph_count >= 4
    table.cell(1, 1).text = "bad\ufffdtext"
    document.save(source)
    checked = FileQualityChecker().validate(artifact, QualityProfile.DOCX, minimum_paragraphs=1)
    assert not checked.passed
    assert [item.code for item in checked.issues] == ["text_corrupted"]


def test_pdf_xlsx_without_table_fails_without_partial_output(tmp_path):
    source = tmp_path / "text.pdf"
    _pdf(source)
    result = converter.convert_pdf_to_office(str(source), "xlsx")
    assert not result.success and result.error_type == "no_reliable_tables"
    assert "可靠表格" in result.error_msg
    assert result.output_path is None
    assert all(event.stage != "completed" for event in result.progress_events)


def test_pdf_xlsx_reliable_table_is_preserved(tmp_path):
    source = tmp_path / "table.pdf"
    _pdf(source, pages=1, table=True)
    result = converter.convert_pdf_to_office(str(source), "xlsx")
    try:
        assert result.success
        workbook = load_workbook(result.output_path, read_only=True)
        try:
            assert list(workbook.active.values)[:2] == [("Item", "Value"), ("Model", "42")]
        finally:
            workbook.close()
    finally:
        converter.cleanup_conversion_output(result.output_path or "")


def test_pptx_capability_warns_before_selection(client, auth_headers):
    headers, _ = auth_headers()
    response = client.get("/file-processing/capabilities?source_format=pdf", headers=headers)
    item = next(item for item in response.json()["capabilities"] if item["capability_id"] == "pdf.to_pptx")
    assert item["editable_output"] is False
    assert "图片" in item["output_description"] and "不可编辑" in item["output_description"]


def test_pdf_worker_reports_real_counts_then_validated_completion(tmp_path):
    source = tmp_path / "two.pdf"
    _pdf(source)
    events = []
    result = pdf_processor.execute_task(FileProcessingRequest(task_type="extract", source_format="pdf",
            source_paths=[str(source)]), progress=events.append)
    assert result.success and result.quality_checked
    page_events = [event for event in events if event.unit == "pages"]
    assert [(event.processed, event.total) for event in page_events] == [(1, 2), (2, 2)]
    assert events[0].total is None
    stages = [event.stage for event in events]
    assert stages[-1] == "completed"
    assert stages.index("validating") < stages.index("completed")


def test_progress_rejects_fake_processed_above_total():
    with pytest.raises(ValidationError):
        FileTaskProgress(stage="converting", processed=2, total=1)
    assert FileTaskProgress(stage="queued").total is None


def test_quality_failure_never_reports_completed(monkeypatch):
    from layers.file_processing.native_text import NativeTextProcessor
    native_text_processor = NativeTextProcessor(lambda *args: "")
    monkeypatch.setattr(native_text_processor, "execute", lambda request: FileProcessingResult(
        success=True, status="SUCCESS"))
    monkeypatch.setattr(native_text_processor, "validate_output", lambda *args: QualityCheckResult(passed=False))
    result = native_text_processor.execute_task(FileProcessingRequest(task_type="write_text",
        target_format="md", output_path="unused"))
    assert not result.success
    assert [event.stage for event in result.progress_events] == ["validating", "failed"]


def test_nested_adapter_does_not_finish_outer_task_early(monkeypatch):
    from layers.file_processing.native_text import NativeTextProcessor
    native_text_processor = NativeTextProcessor(lambda *args: "")
    from layers.file_processing.runner import task_scope
    monkeypatch.setattr(native_text_processor, "execute", lambda request: FileProcessingResult(
        success=True, status="SUCCESS", quality_checked=True))
    with task_scope() as scope:
        result = native_text_processor.execute_task(FileProcessingRequest(task_type="write_text",
            target_format="md", output_path="unused"))
        assert result.success
        assert not any(event.stage == "completed" for event in scope.events)


@pytest.mark.parametrize("status,error_type", [("TIMEOUT", "timeout"), ("CANCELLED", "cancelled"),
                                             ("FAILED", "unsupported_conversion"), ("FAILED", "heavy_task_busy")])
def test_generation_timeout_cancel_never_degrades(tmp_path, monkeypatch, status, error_type):
    monkeypatch.setattr(converter, "convert_file", lambda *args: converter.ConversionResult(
        success=False, status=status, error_type=error_type))
    result = execution.generate_file("# test", "generation-no-fallback", "test", "pdf",
                                    "11111111-1111-1111-1111-111111111111")
    assert not result.success and result.error_type == error_type
    assert result.status.value == status
    assert result.file_id == "" and not result.degradation_code


def test_generated_markdown_degradation_is_explicit_and_has_notice(monkeypatch):
    monkeypatch.setattr(converter, "convert_file", lambda *args: converter.ConversionResult(
        success=False, status="FAILED", error_type="process_failed"))
    result = execution.generate_file("# test", "generation-degraded", "test", "pdf",
                                    "11111111-1111-1111-1111-111111111111")
    assert result.success and result.status == FileProcessingStatus.DEGRADED
    assert result.degradation_code == OFFICE_TO_MARKDOWN
    assert result.user_notice == DEGRADATIONS[OFFICE_TO_MARKDOWN]
    assert result.progress_events[-1].stage == "completed"


def test_raised_file_timeout_cannot_become_markdown_degradation(monkeypatch):
    from layers.file_processing.runner import FileTaskTimeout
    def fail(*args):
        raise FileTaskTimeout("task_deadline")
    monkeypatch.setattr(converter, "convert_file", fail)
    result = execution.generate_file("# test", "generation-raised-timeout", "test", "pdf",
                                    "11111111-1111-1111-1111-111111111111")
    assert result.status == FileProcessingStatus.TIMEOUT
    assert not result.success and not result.file_id and not result.degradation_code


def test_ingest_persisted_progress_sse_protocol(client, user_factory, monkeypatch):
    from tests.test_ingest_batch_progress import _run_task
    monkeypatch.setattr(config, "INGEST_CHUNK_BATCH_SIZE", 2)
    user = user_factory("employee")
    observed = []
    original = task_store.update_task
    def update(task_id, **kwargs):
        original(task_id, **kwargs)
        if kwargs.get("file_progress"):
            observed.append(task_store.get_task(task_id).file_progress)
    monkeypatch.setattr(task_store, "update_task", update)
    task = _run_task(user["user_id"], ["one", "two", "three"], "protocol-doc")
    assert task.status == "done"
    assert task.file_progress.stage == "completed"
    assert task.file_progress.processed == task.file_progress.total == 3
    assert observed[0].stage == "queued"
    assert any(event.stage == "ingesting" and event.processed == 2 and event.total == 3 for event in observed)
    assert [event.stage for event in observed][-2:] == ["validating", "completed"]
    async def collect():
        return [event async for event in main._task_progress_events(task.task_id, user["user_id"])]
    payload = json.loads(asyncio.run(collect())[0].removeprefix("data: ").strip())
    assert payload["file_progress"] == task.file_progress.model_dump(mode="json")


def test_agent_file_progress_is_recorded_and_forwarded(monkeypatch):
    from layers.file_processing.runner import emit_progress
    from layers import planning
    def file_tool(**kwargs):
        emit_progress("converting", 1, 2, "pages")
        return execution.GenerateFileResult(success=False, error_type="process_failed")
    monkeypatch.setattr(execution, "generate_file", file_tool)
    events = []
    state = planning._new_agent_state("file-progress", "file task", "expert")
    state["tool_event_sink"] = events.append
    result = execution.run("generate_file", {}, state=state)
    assert result.status == "error"
    assert state["file_task_progress_events"] == [
        dict(stage="converting", processed=1, total=2, unit="pages")]
    assert any(isinstance(event, FileTaskProgress) for event in events)


def test_sse_file_progress_preserves_unknown_total(monkeypatch):
    from fastapi import BackgroundTasks
    def stream(*args, tool_event_sink=None, **kwargs):
        tool_event_sink(FileTaskProgress(stage="queued"))
        yield main._sse_data({"chunk": "[DONE]"})
    monkeypatch.setattr(main, "_chat_stream_events", stream)
    async def collect():
        return [event async for event in main._chat_stream_events_with_heartbeat(
            main.ChatRequest(session_id="file-progress", message="test"),
            {"user_id": "test"}, BackgroundTasks(), "test-trace", [], [], "test-key")]
    payloads = [json.loads(event.removeprefix("data: ").strip()) for event in asyncio.run(collect())]
    assert payloads[0] == {"file_progress": dict(stage="queued", processed=0, total=None, unit="items")}


def test_repeated_async_cancel_still_waits_for_resource_cleanup():
    import threading
    from layers.file_processing.runner import current_scope
    entered, cleaning, finish = (threading.Event() for _ in range(3))
    def work():
        try:
            entered.set()
            while True:
                current_scope().check()
                finish.wait(.001)
        finally:
            cleaning.set()
            assert finish.wait(2)
    async def scenario():
        task = asyncio.create_task(main._run_file_thread(work))
        assert await asyncio.to_thread(entered.wait, 1)
        task.cancel()
        assert await asyncio.to_thread(cleaning.wait, 1)
        task.cancel()
        await asyncio.sleep(0)
        assert not task.done(), "must wait for cleanup despite repeated disconnect cancellation"
        finish.set()
        with pytest.raises(asyncio.CancelledError):
            await task
    asyncio.run(scenario())


def test_pdf_split_api_exposes_real_page_progress(client, auth_headers, tmp_path):
    headers, _ = auth_headers()
    source = tmp_path / "split.pdf"
    _pdf(source)
    response = client.post("/tools/pdf/split", headers=headers,
        files={"file": ("split.pdf", source.read_bytes(), "application/pdf")})
    assert response.status_code == 200
    events = response.json()["progress_events"]
    assert [(item["processed"], item["total"]) for item in events if item["unit"] == "pages"] == [(1, 2), (2, 2)]
    assert events[-1]["stage"] == "completed"


def test_http_disconnect_cancels_file_worker_without_external_task_cancel():
    import threading
    from layers.file_processing.runner import current_scope, FileTaskCancelled
    entered, disconnected, cleaned = (threading.Event() for _ in range(3))
    class Request:
        async def is_disconnected(self):
            return disconnected.is_set()
    def work():
        try:
            entered.set()
            while True:
                current_scope().check()
                cleaned.wait(.001)
        finally:
            cleaned.set()
    async def scenario():
        task = asyncio.create_task(main._run_file_thread(work, http_request=Request()))
        assert await asyncio.to_thread(entered.wait, 1)
        disconnected.set()
        with pytest.raises(FileTaskCancelled):
            await asyncio.wait_for(task, 1)
        assert cleaned.is_set()
    asyncio.run(scenario())
