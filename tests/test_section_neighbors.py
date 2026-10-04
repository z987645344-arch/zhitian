"""同节邻段只扩展最终生成资料，真实Chroma + 确定性模型桩。"""
import json
from unittest.mock import Mock

import numpy as np
import pytest
from docx import Document

import config
import main
from layers import auth, document_loader, document_sections, execution, memory, planning
from tests.test_source_policy import state, response, classified


def paths(*parts):
    return json.dumps([list(parts)], ensure_ascii=False)


def context(doc="a", indices=(1,), allowed=None):
    return execution.DocumentAnswerContext(query="问题", tier="expert", allowed_doc_ids=allowed or [doc],
        candidates=[execution.DocumentAnswerCandidate(content=f"原资料{i}", source="测试.md",
            score=.6, doc_id=doc, chunk_index=i) for i in indices])


def seed(doc="a", sections=None):
    memory.save_document("测试.md", [f"原资料{i}" for i in range(5)], doc,
        chunk_section_paths=sections if sections is not None else [paths("标题", "节一")] * 5)


def test_markdown_sections_cross_boundaries_and_skip_code_fences():
    text = "# 标题\n\n## 第一节\n甲甲甲。\n```md\n## 假标题\n```\n乙乙乙。\n## 第二节\n丙丙丙。"
    chunks = document_loader.chunk_text(text, chunk_size=38)
    before = list(chunks)
    fields = document_sections.chunk_section_paths(text, chunks, source_name="测试.md")
    assert chunks == before
    assert all(fields)
    assert all("假标题" not in value for value in fields)
    all_paths = set().union(*(document_sections.read_section_paths(value) for value in fields))
    assert ("标题", "第一节") in all_paths and ("标题", "第二节") in all_paths
    crossing = document_sections.chunk_section_paths(text, [text], source_name="测试.md")
    assert {("标题", "第一节"), ("标题", "第二节")} <= document_sections.read_section_paths(crossing[0])
    assert document_sections.chunk_section_paths("无标题正文", ["无标题正文"], source_name="测试.txt") == [""]
    assert document_sections.chunk_section_paths(text, ["不匹配正文"], source_name="测试.md") == [""]
    multiple = "# 文档标题\n简介\n# 第一章\n## 一节\n正文"
    assert ("文档标题", "第一章", "一节") in document_sections.read_section_paths(
        document_sections.chunk_section_paths(multiple, [multiple], source_name="测试.md")[0])


def test_docx_heading_styles_metadata_without_changing_text(tmp_path):
    doc = Document()
    doc.add_heading("文档标题", level=0)
    doc.add_heading("大节", level=1)
    doc.add_paragraph("普通正文")
    doc.add_heading("小节", level=2)
    doc.add_paragraph("第二段正文")
    target = tmp_path / "标题.docx"
    doc.save(target)
    text = document_loader.load_document(str(target))
    chunks = document_loader.chunk_text(text)
    fields = document_sections.chunk_section_paths(text, chunks, source_path=str(target))
    assert document_loader.load_document(str(target)) == text
    assert ("文档标题", "大节", "小节") in document_sections.read_section_paths(fields[0])


def test_optional_metadata_keeps_documents_indices_and_vectors_identical():
    text = "# 标题\n## 第一节\n" + "甲乙丙丁。" * 160
    chunks = document_loader.chunk_text(text)
    fields = document_sections.chunk_section_paths(text, chunks, source_name="测试.md")
    memory.save_document("测试.md", chunks, "old")
    memory.save_document("测试.md", chunks, "new", chunk_section_paths=fields)
    collection = memory._get_document_collection()
    old = collection.get(where={"doc_id": "old"}, include=["documents", "metadatas", "embeddings"])
    new = collection.get(where={"doc_id": "new"}, include=["documents", "metadatas", "embeddings"])
    sort = lambda rows: sorted(zip(rows["metadatas"], rows["documents"], rows["embeddings"]), key=lambda row: row[0]["chunk_index"])
    for (om, ot, ov), (nm, nt, nv) in zip(sort(old), sort(new)):
        assert ot == nt
        assert om["chunk_index"] == nm["chunk_index"]
        np.testing.assert_array_equal(ov, nv)
        assert "section_paths" not in om and nm["section_paths"]
    # 原有读取模型只投影既有字段，新字段是Chroma可接受的JSON字符串标量。
    clean = [{k: v for k, v in item.items() if k != "section_paths"} for item in new["metadatas"]]
    payload = dict(documents=[new["documents"]], distances=[[.2] * len(chunks)])
    assert memory._build_document_search_results({**payload, "metadatas": [clean]}, None, None) == memory._build_document_search_results({**payload, "metadatas": [new["metadatas"]]}, None, None)


def test_neighbors_same_doc_shared_section_adjacent_and_deduplicated():
    seed(sections=[paths("标题", "节一"), paths("标题", "节一"), paths("标题", "节一"),
                   paths("标题", "节二"), paths("标题", "节二")])
    seed("outside")
    result = memory.get_document_section_neighbors([("a", 1), ("outside", 1)], ["a"], 2)
    assert [(item.doc_id, item.chunk_index) for item in result] == [("a", 0), ("a", 2)]
    assert all("score" not in item.model_dump() for item in result)
    assert memory.get_document_section_neighbors([("a", 2)], ["a"], 2)[0].chunk_index == 1
    assert not memory.get_document_section_neighbors([("a", 1)], [], 2)
    assert not memory.get_document_section_neighbors([("a", 1)], ["a"], 0)
    assert [item.chunk_index for item in memory.get_document_section_neighbors([("a", 1), ("a", 0)], ["a"], 2)] == [2]


@pytest.mark.parametrize("field", ["", "[]", "{broken", '["not a path"]'])
def test_missing_or_invalid_sections_never_expand(field):
    seed(sections=[field] * 5)
    assert memory.get_document_section_neighbors([("a", 1)], ["a"], 2) == []


def test_request_limit_idempotence_disable_and_current_verified_scope(monkeypatch):
    seed("a")
    seed("b")
    monkeypatch.setattr(auth, "get_verified_doc_ids", lambda: ["a", "b"])
    request = {}
    first = context()
    execution.prepare_document_answer_context(first, request)
    assert [item.chunk_index for item in first.candidates] == [1, 0, 2]
    assert request["section_neighbor_count"] == 2
    execution.prepare_document_answer_context(first, request)
    second = context("b")
    execution.prepare_document_answer_context(second, request)
    assert len(first.candidates) == 3 and len(second.candidates) == 1
    # 下一次请求独立计数；刚撤销核验及不在最初范围的文档不能补取。
    monkeypatch.setattr(auth, "get_verified_doc_ids", lambda: ["b"])
    assert len(execution.prepare_document_answer_context(context(), {}).candidates) == 1
    monkeypatch.setattr(config, "RAG_SECTION_NEIGHBOR_MAX", 0)
    assert len(execution.prepare_document_answer_context(context("b"), {}).candidates) == 1


@pytest.mark.parametrize("stream", [False, True])
def test_expert_final_context_expanded_only_at_respond_and_citations_agree(monkeypatch, stream):
    seed()
    monkeypatch.setattr(auth, "get_verified_doc_ids", lambda: ["a"])
    request = state(source="internal", evidence="weak")
    ctx = context()
    result = execution.ToolResult(tool="search_documents", status="success", data="原检索数据",
        citations=execution.document_answer_citations(ctx), document_answer_context=ctx,
        metadata={"document_answer_deferred": True})
    request.update(intent="document", results=[result], stream_document_answer=stream)
    assert len(ctx.candidates) == 1
    generator = Mock(return_value=iter(["成品回答"]))
    monkeypatch.setattr(execution, "_answer_from_documents", generator)
    planning.respond_node(request)
    assert [item.chunk_index for item in ctx.candidates] == [1, 0, 2]
    assert [item.chunk_index for item in request["citations"]] == [1, 0, 2]
    serialized = main._serialize_citations(request["citations"])
    assert "score" not in serialized[1] and "score" not in serialized[2]
    if stream:
        generator.assert_not_called()
        assert main._streamable_document_answer_context(request) is ctx
    else:
        assert generator.call_args.args[0] is ctx


def test_fast_expand_selected_only_evidence_filter_input_unchanged(monkeypatch):
    seed()
    monkeypatch.setattr(auth, "get_verified_doc_ids", lambda: ["a"])
    request = state(mode="fast", source="internal", evidence="hit")
    ctx = context(indices=(1, 4))
    ctx.candidates[1].content = "未选资料4"
    result = execution.ToolResult(tool="search_documents", status="success", data="[1] 原资料1\n\n[2] 未选资料4",
        citations=execution.document_answer_citations(ctx), document_answer_context=ctx)
    before = planning._build_fast_evidence_messages(request, result)
    text, citations = execution.prepare_fast_document_evidence(result, "[1] 原资料1", [result.citations[0]], request)
    assert before == planning._build_fast_evidence_messages(request, result)
    assert len(ctx.candidates) == 2
    assert "原资料0" in text and "原资料2" in text
    assert "未选资料4" not in text and "原资料3" not in text
    assert [item.chunk_index for item in citations] == [1, 0, 2]


def test_fast_real_node_selection_before_expansion(monkeypatch):
    seed()
    monkeypatch.setattr(auth, "get_verified_doc_ids", lambda: ["a"])
    ctx = context(indices=(1, 4))
    result = execution.ToolResult(tool="search_documents", status="success", data="[1] 原资料1\n\n[2] 原资料4",
        citations=execution.document_answer_citations(ctx), document_answer_context=ctx)
    observed = {}
    def model(messages, **kwargs):
        stage = kwargs["stage"]
        observed[stage] = messages
        if stage == "fast_tool_selection":
            return response(classification=classified(source="internal"))
        if stage == "fast_evidence_filter":
            return {"choices": [{"message": {"content": json.dumps({"evidence_sufficient": True,
                "used_candidate_ids": [1], "reason": "sufficient"})}}]}
        return {"choices": [{"message": {"content": "成品答案"}}]}
    monkeypatch.setattr(planning.llm_provider, "chat_completion", model)
    monkeypatch.setattr(planning.mcp_client, "call_tool", lambda *_a, **_kw: result)
    final = planning.run_graph_state("neighbor-fast", "问题", "fast")
    assert "原资料0" not in str(observed["fast_evidence_filter"])
    assert "原资料0" in str(observed["fast_result_generation"])
    assert "原资料2" in str(observed["fast_result_generation"])
    assert "原资料4" not in str(observed["fast_result_generation"])
    assert [item.chunk_index for item in final["citations"]] == [1, 0, 2]


def test_metadata_count_mismatch_rejected_before_writing():
    with pytest.raises(ValueError, match="小节元数据"):
        memory.save_document("测试.md", ["正文"], "mismatch", chunk_section_paths=[])
    assert memory.count_document_chunks("mismatch") == 0


@pytest.mark.parametrize("route", ["upload", "manual"])
def test_real_ingest_endpoints_store_optional_sections(client, auth_headers, route):
    from tests.conftest import grant_work_organization
    headers, user = auth_headers("employee")
    organization_id = grant_work_organization(user["user_id"])
    text = "# 测试文档\n## 功能\n" + "这是用于测试的普通资料。" * 80
    if route == "upload":
        reply = client.post("/documents/upload", headers=headers,
            files={"file": ("test-sections.md", text.encode(), "text/markdown")},
            data={"organization_id": organization_id})
    else:
        reply = client.post("/knowledge/input", headers=headers,
            json={"content": text, "title": "测试文档", "organization_id": organization_id})
    assert reply.status_code == 200, reply.text
    doc_id = reply.json()["doc_id"]
    rows = memory._get_document_collection().get(where={"doc_id": doc_id}, include=["metadatas"])
    assert rows["metadatas"] and all(metadata.get("section_paths") for metadata in rows["metadatas"])


def test_neighbor_fetch_error_is_logged_and_original_materials_retained(monkeypatch, caplog):
    monkeypatch.setattr(auth, "get_verified_doc_ids", lambda: ["a"])
    monkeypatch.setattr(memory, "get_document_section_neighbors", Mock(side_effect=OSError("hidden data")))
    original = context()
    execution.prepare_document_answer_context(original, {})
    assert len(original.candidates) == 1
    assert "error_type=OSError" in caplog.text
    assert "hidden data" not in caplog.text
